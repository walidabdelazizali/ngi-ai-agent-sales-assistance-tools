"""Extract structured data from .docx files."""

import json
from pathlib import Path
from typing import Any

from docx import Document

from src.config.settings import INPUT_DIR, OUTPUT_DIR
from src.extractors.schema import SCHEMA_VERSION
from src.normalizers.extraction_normalizer import normalize_extraction
from src.validators.extraction_validator import validate_or_raise

SUPPORTED_EXTENSIONS = {".docx"}


def extract_paragraphs(doc: Document) -> list[str]:
    """Return non-empty paragraph texts."""
    return [p.text for p in doc.paragraphs if p.text.strip()]


def extract_tables(doc: Document) -> list[list[list[str]]]:
    """Return tables as list of rows, each row a list of cell texts."""
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return tables


def extract_docx(file_path: Path) -> dict[str, Any]:
    """Extract paragraphs and tables from a single .docx file."""
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {file_path.suffix}")

    doc = Document(str(file_path))
    paragraphs = extract_paragraphs(doc)
    tables = extract_tables(doc)

    return {
        "schema_version": SCHEMA_VERSION,
        "source_filename": file_path.name,
        "source_path": str(file_path),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
    }


def run_ingestion() -> list[Path]:
    """Process all .docx files in INPUT_DIR, write JSON to OUTPUT_DIR.

    Returns list of output file paths written.
    """
    if not INPUT_DIR.is_dir():
        raise FileNotFoundError(f"Input directory missing: {INPUT_DIR}")

    docx_files = sorted(INPUT_DIR.glob("*.docx"))
    if not docx_files:
        print(f"No .docx files found in {INPUT_DIR}")
        return []

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []
    for docx_path in docx_files:
        result = extract_docx(docx_path)
        result = normalize_extraction(result)
        validate_or_raise(result)
        out_path = OUTPUT_DIR / f"{docx_path.stem}.json"
        out_path.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"  {docx_path.name} -> {out_path.name}")
        written.append(out_path)

    return written
