import openpyxl
import pandas as pd
from pathlib import Path
import re

def snake_case(s):
    s = s.strip().replace("\n", " ")
    s = re.sub(r"[^\w\s]", "", s)
    s = re.sub(r"[\s\-]+", "_", s)
    return s.lower()

def find_header_row(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and any(cell and isinstance(cell, str) for cell in row):
            # Heuristic: header row has at least 2 non-empty string cells
            non_empty = [cell for cell in row if cell and isinstance(cell, str)]
            if len(non_empty) >= 2:
                return i, [cell.strip() if isinstance(cell, str) else cell for cell in row]
    raise ValueError(f"No header row found in sheet {ws.title}")

def parse_sheet(wb, sheet_name, out_path):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet '{sheet_name}' not found, skipping.")
        return None
    ws = wb[sheet_name]
    header_row_idx, headers = find_header_row(ws)
    headers_snake = [snake_case(h) if h else f"col_{i+1}" for i, h in enumerate(headers)]
    data = []
    for i, row in enumerate(ws.iter_rows(min_row=header_row_idx+1, values_only=True), start=header_row_idx+1):
        if not any(row):
            continue
        cleaned = [(str(cell).strip() if cell is not None else "") for cell in row]
        data.append(cleaned)
    df = pd.DataFrame(data, columns=headers_snake)
    df.to_csv(out_path, index=False, encoding="utf-8-sig")
    print(f"Wrote {len(df)} rows to {out_path}")
    return headers, out_path, len(df)

def parse_network_workbook(xlsx_path, output_dir):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True, keep_links=False)
    sheets = [
        ("Network List", "network_list_normalized.csv"),
        ("Dental Network", "dental_network_normalized.csv"),
        ("Optical Network", "optical_network_normalized.csv"),
        ("Network Additions", "network_additions_normalized.csv"),
        ("Network Deletions", "network_deletions_normalized.csv"),
        ("Network Recategorizations", "network_recategorizations_normalized.csv"),
    ]
    results = []
    for sheet, fname in sheets:
        out_path = Path(output_dir) / fname
        try:
            res = parse_sheet(wb, sheet, out_path)
            if res:
                headers, path, count = res
                results.append({
                    "sheet": sheet,
                    "headers": headers,
                    "output": str(path),
                    "row_count": count
                })
        except Exception as e:
            print(f"Error parsing {sheet}: {e}")
    return results
