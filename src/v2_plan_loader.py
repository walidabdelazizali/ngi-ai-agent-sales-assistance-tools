import os
from docx import Document

# Map plan names to DOCX filenames (update with actual filenames/paths)

# V2: Only allow plan loading from input_docs/ with explicit mapping to actual filenames
PLAN_DOCX_MAP = {
    "Remedy 02": "input_docs/HN-REMEDY-2.docx",
    "Remedy 03": "input_docs/HN-REMEDY-3.docx",
    "Remedy 04": "input_docs/HN-REMEDY 4.docx",
    "Remedy 05": "input_docs/HN-REMEDY 5.docx",
    "Remedy 06": "input_docs/HN-REMEDY 6.docx",
}

def load_clean_plan(plan_name: str) -> dict:
    """
    Load and extract raw data from the original Remedy DOCX file for the given plan.
    Returns a dict with paragraphs, tables, and raw text.
    """
    if plan_name not in PLAN_DOCX_MAP:
        raise ValueError(f"Unsupported plan name: {plan_name}")
    docx_path = PLAN_DOCX_MAP[plan_name]
    # Enforce input_docs/ only
    if not docx_path.startswith("input_docs/"):
        raise ValueError(f"Plan source path must be under input_docs/: {docx_path}")
    if not os.path.isfile(docx_path):
        raise ValueError(f"DOCX file not found for {plan_name}: {docx_path}")
    doc = Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        tables.append(table_data)
    raw_text = "\n".join(paragraphs + ["\n".join(["\t".join(row) for row in tbl]) for tbl in tables])
    # Minimal patch: extract network_name using same logic as v2_plan_normalizer
    def norm(val):
        if val is None:
            return None
        v = str(val).strip()
        if v.lower() in ("", "none", "null", "not available", "n/a"):
            return None
        if v.upper() == "NIL":
            return "Nil"
        return v

    def extract_value(val, label=None):
        if val is None:
            return None
        v = str(val).strip()
        if label and v.lower().startswith(label.lower()):
            v = v[len(label):].lstrip(': .')
        if ':' in v:
            v = v.split(':', 1)[1].strip()
        return norm(v)

    def find_table_value(label, tables, value_only=False):
        for table in tables:
            for row in table:
                if len(row) >= 2 and label.lower() in row[0].lower():
                    return extract_value(row[1]) if value_only else norm(row[1])
        return None

    def find_paragraph_value(keyword, paragraphs, value_only=False):
        for p in paragraphs:
            if keyword.lower() in p.lower():
                return extract_value(p, keyword) if value_only else norm(p)
        return None

    network_name = find_table_value("network", tables, value_only=True) or find_paragraph_value("network", paragraphs, value_only=True)

    result = {
        "plan_name": plan_name,
        "source_path": docx_path,
        "paragraphs": paragraphs,
        "tables": tables,
        "raw_text": raw_text,
    }
    if network_name:
        result["network_name"] = network_name
    return result

if __name__ == "__main__":
    from src.v2_plan_normalizer import normalize_clean_plan
    import pprint
    for plan in PLAN_DOCX_MAP:
        try:
            raw = load_clean_plan(plan)
            print(f"Plan: {plan}\n  source_path: {raw['source_path']}")
            norm = normalize_clean_plan(raw)
            print(f"  paragraph count: {len(raw['paragraphs'])}")
            print(f"  table count: {len(raw['tables'])}")
            print("  normalized output:")
            pprint.pprint(norm)
        except Exception as e:
            print(f"Plan: {plan}\n  ERROR: {e}")
