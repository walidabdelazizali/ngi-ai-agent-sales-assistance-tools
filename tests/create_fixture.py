"""Generate a minimal .docx fixture for tests."""

from pathlib import Path

from docx import Document


def create_fixture(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


if __name__ == "__main__":
    create_fixture(Path("tests/fixtures/sample.docx"))
    print("Fixture created.")
