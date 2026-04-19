import pytest
from pathlib import Path
import openpyxl
from src.parsers import network_workbook_parser
import pandas as pd

def test_workbook_opens():
    path = Path("data/networks/raw/uae_network_feb_2026_source.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    assert wb is not None

def test_detected_sheets():
    path = Path("data/networks/raw/uae_network_feb_2026_source.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    expected = [
        "Network List",
        "Dental Network",
        "Optical Network",
        "Network Additions",
        "Network Deletions",
        "Network Recategorizations"
    ]
    found = [s for s in expected if s in wb.sheetnames]
    assert found, "No expected sheets found"

def test_csv_outputs(tmp_path):
    src = Path("data/networks/raw/uae_network_feb_2026_source.xlsx")
    outdir = tmp_path
    results = network_workbook_parser.parse_network_workbook(src, outdir)
    for r in results:
        out = Path(r["output"])
        assert out.exists()
        df = pd.read_csv(out)
        assert len(df) == r["row_count"]

def test_network_list_core_columns(tmp_path):
    src = Path("data/networks/raw/uae_network_feb_2026_source.xlsx")
    outdir = tmp_path
    results = network_workbook_parser.parse_network_workbook(src, outdir)
    for r in results:
        if r["sheet"] == "Network List":
            df = pd.read_csv(r["output"])
            cols = set(df.columns)
            # Example core columns, adjust as needed
            assert "provider_name" in cols or "facility_name" in cols
            assert "city" in cols or "location" in cols
            break
